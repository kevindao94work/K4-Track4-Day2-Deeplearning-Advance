"""Workbook từ kết quả thật; openpyxl fallback khi artifact-tool không có trên host."""
from __future__ import annotations
import json, math, textwrap
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule
from PIL import Image, ImageDraw, ImageFont
from paths import SUB
from train import write_json

SHEETS=['Summary','Backbones','Training','Inference','Final','PerClass','Latency']

def source_tables():
    frames={n:pd.read_csv(SUB/'tables'/f'{n}.csv') for n in SHEETS if n!='Summary'}
    for name,late in [('Backbones',['pretrained_revision']),('Inference',['checkpoint_sha256'])]:
        frames[name]=frames[name][[c for c in frames[name] if c not in late]+late]
    first=['configuration','gpu','dtype','batch','bn_fusion','p50','p95','p99','images_per_s']
    frames['Latency']=frames['Latency'][first+[c for c in frames['Latency'] if c not in first]]
    frames['Final']['row_type']='Seed riêng'
    aggregate=[]
    for eid,g in frames['Final'].groupby('exp_id',sort=False):
        row={'exp_id':eid,'complete_configuration':g.complete_configuration.iloc[0],
             'seed':'0/1/2','row_type':'Tổng hợp ba seed',
             'mean_std_summary':g.mean_std_summary.iloc[0]}
        for c in ['macro_f1_val','macro_f1_test','top1_test','ece_test','nll_test']:
            row[c]=float(g[c].mean());row[c+'_std']=float(g[c].std(ddof=1))
        row['macro_f1_test_mean']=row['macro_f1_test']
        aggregate.append(row)
    frames['Final']=pd.concat([frames['Final'],pd.DataFrame(aggregate)],ignore_index=True)
    rank=[]
    for phase in ['Backbones','Training','Inference']:
        for _,r in frames[phase].iterrows():
            p95=r.get('latency_p95_batch1_ms',None)
            p50=r.get('latency_p50_batch1_ms',r.get('batch1_latency_ms',None))
            rank.append({'exp_id':r.exp_id,'giai_doan':phase,'so_seed':1,
                'macro_f1_val':r.macro_f1_val,'p50_batch1_ms':p50,'p95_batch1_ms':p95,
                'chi_phi_so_I00':r.get('relative_cost_vs_I00',None),
                'ghi_chu':('Độ trễ sơ bộ 5 mẫu' if phase=='Backbones' else
                            'Chưa benchmark riêng' if phase=='Training' else '50 mẫu, checkpoint sàng lọc')})
    for row in aggregate:
        lat=frames['Latency'];lat=lat[(lat.configuration==row['exp_id']+'_seed0')&(lat.batch==1)].iloc[0]
        rank.append({'exp_id':row['exp_id'],'giai_doan':'Chung kết val','so_seed':3,
            'macro_f1_val':row['macro_f1_val'],'p50_batch1_ms':lat.p50,'p95_batch1_ms':lat.p95,
            'chi_phi_so_I00':None,'ghi_chu':'Mean 3 seed; latency checkpoint seed 0; không chọn lại'})
    frames['Summary']=pd.DataFrame(rank).sort_values('macro_f1_val',ascending=False,kind='stable').head(10).reset_index(drop=True)
    frames['Summary'].insert(0,'hang_val',range(1,11))
    frames['Summary'].to_csv(SUB/'tables/Summary.csv',index=False)
    return frames

LABELS={'exp_id':'Mã thí nghiệm','backbone':'Backbone','pretrained_tag':'Tag trọng số',
 'pretrained_revision':'Revision nguồn','parameters_m':'Tham số (M)','gmac':'GMAC đo',
 'resolution':'Độ phân giải','epochs':'Epoch','seed':'Seed','macro_f1_val':'Macro-F1 val',
 'top1_val':'Top-1 val','training_seconds_per_epoch':'Train/epoch (s)',
 'seconds_per_epoch_train_val':'Train+val/epoch (s)','batch1_latency_ms':'p50 sơ bộ batch 1 (ms)',
 'best_epoch':'Epoch tốt nhất','notes':'Ghi chú','changed_axis':'Trục thay đổi',
 'difference_from_T00':'Thay đổi so T00','delta_vs_T00':'Δ macro-F1 so T00',
 'chinee_apple_f1_val':'F1 Chinee val','snake_weed_f1_val':'F1 Snake val','ece_val':'ECE val',
 'method':'Phương pháp','model_checkpoint':'Checkpoint','checkpoint_sha256':'SHA-256 checkpoint',
 'K':'Số view K','input_size':'Đầu vào tiền xử lý','temperature':'Nhiệt độ T',
 'nll_val':'NLL val','latency_p50_batch1_ms':'p50 batch 1 (ms)',
 'latency_p95_batch1_ms':'p95 batch 1 (ms)','latency_p99_batch1_ms':'p99 batch 1 (ms)',
 'throughput_images_s':'Thông lượng batch 32 (ảnh/s)','relative_cost_vs_I00':'Chi phí p50 so I00',
 'complete_configuration':'Cấu hình đầy đủ','macro_f1_test':'Macro-F1 test','top1_test':'Top-1 test',
 'ece_test':'ECE test','nll_test':'NLL test','temperature_from_val':'T khớp val',
 'mean_std_summary':'Mean ± sample std test','macro_f1_test_mean':'Mean macro-F1 test',
 'macro_f1_test_std':'Std macro-F1 test','macro_f1_val_std':'Std macro-F1 val',
 'top1_test_std':'Std top-1 test','ece_test_std':'Std ECE test','nll_test_std':'Std NLL test',
 'row_type':'Loại dòng','class_label':'Nhãn lớp','class':'Lớp','test_count':'Số ảnh test/seed',
 'precision':'Precision test mean','recall':'Recall test mean','f1':'F1 test mean',
 'precision_std':'Std precision','recall_std':'Std recall','f1_std':'Std F1',
 'p50':'p50 (ms)','p95':'p95 (ms)','p99':'p99 (ms)','mean':'Mean (ms)',
 'iterations':'Số mẫu thời gian','warmups':'Warmup','gpu':'Thiết bị','dtype':'Kiểu số',
 'batch':'Batch','bn_fused':'Đã gộp BN','torch_version':'Phiên bản torch',
 'images_per_second':'Thông lượng (ảnh/s)','preprocessing_included':'Có tiền xử lý CPU',
 'timing_scope':'Phạm vi đo','hang_val':'Hạng val','giai_doan':'Giai đoạn','so_seed':'Số seed',
 'p50_batch1_ms':'p50 batch 1 (ms)','p95_batch1_ms':'p95 batch 1 (ms)',
 'chi_phi_so_I00':'Chi phí p50 so I00','ghi_chu':'Ghi chú','configuration':'Mã cấu hình',
 'condition':'Phạm vi đo','images_per_s':'Thông lượng (ảnh/s)','bn_fusion':'Đã gộp BN',
 'torch':'Phiên bản torch','img_size':'Độ phân giải','n':'Số mẫu thời gian'}

def value(v):
    if pd.isna(v):return None
    if isinstance(v,np.generic):return v.item()
    return v

def render_saved_sheet(ws, output):
    """Preview đọc file XLSX đã lưu; bảng rộng tách thành từng cụm cột, không bỏ cột."""
    font_path=Path('/System/Library/Fonts/Supplemental/Arial.ttf')
    if not font_path.is_file():
        candidates=list(Path('/usr/share/fonts').rglob('DejaVuSans.ttf'))
        font_path=candidates[0] if candidates else None
    font=ImageFont.truetype(str(font_path),15) if font_path else ImageFont.load_default()
    bold=font
    def wrap_pixels(s,w):
        result=[];line=''
        for char in s:
            if font.getlength(line+char)>w-12:
                result.append(line);line=''
            line+=char
        result.append(line)
        return result
    for start in range(1,ws.max_column+1,6):
        end=min(start+5,ws.max_column)
        widths=[min(380,max(145,int(ws.column_dimensions[get_column_letter(c)].width*8))) for c in range(start,end+1)]
        body=list(range(5,ws.max_row+1))
        row_heights=[];lines=[]
        for r in body:
            vals=[]
            for c,w in zip(range(start,end+1),widths):
                cell=ws.cell(r,c);v=cell.value
                s='' if v is None else f'{v:.2f}' if isinstance(v,float) and cell.number_format=='0.00' else f'{v:.4f}' if isinstance(v,float) else str(v)
                if r==5:s=LABELS.get(s,s)
                vals.append(wrap_pixels(s,w))
            lines.append(vals);row_heights.append(max(32,max(len(x) for x in vals)*19+12))
        canvas=Image.new('RGB',(sum(widths)+40,sum(row_heights)+90),'white');d=ImageDraw.Draw(canvas)
        d.text((20,12),str(ws['A2'].value),fill='#172B4D',font=bold)
        d.text((20,38),f'{ws.title} · cột {start}–{end} · preview từ XLSX đã lưu',fill='#536278',font=font)
        y=75
        for r,vals,h in zip(body,lines,row_heights):
            x=20
            for col,w,txt in zip(range(start,end+1),widths,vals):
                fill='#203864' if r==5 else '#F3F6FA' if r%2==0 else 'white'
                d.rectangle((x,y,x+w,y+h),fill=fill)
                d.multiline_text((x+6,y+6),'\n'.join(txt),fill='white' if r==5 else '#172B4D',font=font,spacing=3)
                x+=w
            y+=h
        canvas.save(output/f'workbook_{ws.title}_{start}.png')

def build_workbook():
    if not json.loads((SUB/'evidence/stage11.json').read_text())['passed']:raise ValueError('Thiếu kết quả test chính thức')
    frames=source_tables();wb=Workbook();wb.remove(wb.active)
    for name in SHEETS:
        df=frames[name];ws=wb.create_sheet(name);ws.sheet_view.showGridLines=False
        ws['A2']=f'Lab Day 2 · {name} · Đào Gia Bảo · 2A202602793'
        ws['A2'].font=Font(name='Arial',size=14,bold=True,color='203864')
        ws['A3']=('Top 10 theo validation; không chọn lại bằng test. Độ trễ không gồm tiền xử lý CPU.' if name=='Summary'
                   else f'Nguồn: tables/{name}.csv; tỷ lệ 0–1; std mẫu ddof=1.' )
        ws['A3'].font=Font(name='Arial',size=10,italic=True,color='536278')
        for i,c in enumerate(df.columns,1):
            cell=ws.cell(5,i,LABELS.get(c,c));cell.fill=PatternFill('solid',fgColor='203864')
            cell.font=Font(name='Arial',size=10,bold=True,color='FFFFFF');cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
            cell.border=Border(bottom=Side(style='thin',color='FFFFFF'))
        ws.row_dimensions[5].height=42
        for i,row in enumerate(df.itertuples(index=False,name=None),6):
            for j,v in enumerate(row,1):
                cell=ws.cell(i,j,value(v));cell.font=Font(name='Arial',size=10,color='172B4D')
                cell.alignment=Alignment(horizontal='right' if isinstance(cell.value,(int,float)) else 'left',vertical='center',wrap_text=True)
                if i%2==0:cell.fill=PatternFill('solid',fgColor='F3F6FA')
                if isinstance(cell.value,float):cell.number_format='0.0000'
                if isinstance(cell.value,int):cell.number_format='#,##0'
                col=df.columns[j-1]
                if any(k in col for k in ['_ms','p50','p95','p99','seconds','throughput','images_per_second','images_per_s']) or col=='mean':
                    if isinstance(cell.value,(int,float)):cell.number_format='0.00'
                if col=='delta_vs_T00':cell.number_format='+0.0000;-0.0000;0.0000'
            ws.row_dimensions[i].height=48 if name=='Final' else 36
        for j,c in enumerate(df.columns,1):
            # Long configuration/hash cells remain fully readable through wrapping.
            width=48 if c in ['complete_configuration','timing_scope','condition','notes'] else 42 if 'sha256' in c or c=='pretrained_revision' else 30 if c in ['method','difference_from_T00','mean_std_summary','ghi_chu','backbone','model_checkpoint'] else 22 if c in ['exp_id','configuration'] else 17
            ws.column_dimensions[get_column_letter(j)].width=width
        if name=='Final':
            for r in range(6,ws.max_row+1):ws.row_dimensions[r].height=130
        if name=='Latency':
            for r in range(6,ws.max_row+1):ws.row_dimensions[r].height=66
        for r in range(6,ws.max_row+1):
            needed=max(len(textwrap.wrap(str(ws.cell(r,c).value or ''),width=max(10,int(ws.column_dimensions[get_column_letter(c)].width)-2))) for c in range(1,ws.max_column+1))
            ws.row_dimensions[r].height=max(ws.row_dimensions[r].height,needed*13+8)
        ws.freeze_panes='C6' if name!='Summary' else 'B6'
        ws.auto_filter.ref=f'A5:{get_column_letter(ws.max_column)}{ws.max_row}'
        ws.print_options.horizontalCentered=True;ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=1 if name=='Summary' else 0
        ws.print_title_rows='1:5';ws.print_area=f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'
        if 'macro_f1_val' in df:
            col=get_column_letter(df.columns.get_loc('macro_f1_val')+1)
            ws.conditional_formatting.add(f'A6:{get_column_letter(ws.max_column)}{ws.max_row}',
                FormulaRule(formula=[f'${col}6=MAX(${col}$6:${col}${ws.max_row})'],
                            fill=PatternFill('solid',fgColor='E2EFDA'),font=Font(bold=True,color='235A28')))
        ws.sheet_properties.tabColor='203864' if name in ['Summary','Final'] else '7593B7'
    path=SUB/'results.xlsx';wb.save(path)
    saved=load_workbook(path,data_only=True);checks={}
    if saved.sheetnames!=SHEETS:raise AssertionError('Sai sheet')
    for name,df in frames.items():
        ws=saved[name]
        if ws.max_row!=len(df)+5 or ws.max_column!=len(df.columns):raise AssertionError(name)
        for i,row in enumerate(df.itertuples(index=False,name=None),6):
            for j,v in enumerate(row,1):
                expected=value(v);actual=ws.cell(i,j).value
                if isinstance(expected,float):
                    if not math.isclose(actual,expected,rel_tol=1e-14,abs_tol=1e-14):raise AssertionError((name,i,j))
                elif actual!=expected:raise AssertionError((name,i,j,actual,expected))
        if any(c.data_type=='e' for row in ws for c in row):raise AssertionError('Ô lỗi')
        checks[name]={'rows':len(df),'columns':len(df.columns),'all_values_match_sources':True}
        render_saved_sheet(ws,SUB/'evidence')
    write_json(SUB/'evidence/stage12.json',{'passed':True,'sheets':checks,'output':str(path),
        'authoring':'openpyxl fallback: artifact-tool/loader và cache runtime không có trên host macOS',
        'formulas':0,'data':'Số đo cố định; bảng tổng hợp tính từ CSV bằng pandas/numpy, không cần recalculation Excel',
        'visual_previews':'evidence/workbook_<sheet>_<start_column>.png; tất cả cột được chia cụm',
        'selection':'Highlight macro-F1 validation cao nhất, không chọn theo test'})
    print('Đã tạo và đối chiếu từng ô của',path)
    return path

if __name__=='__main__':build_workbook()
