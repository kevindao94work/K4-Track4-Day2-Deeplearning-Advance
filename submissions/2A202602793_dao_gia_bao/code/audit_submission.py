"""Kiểm toán hồ sơ đã nộp bằng CSV/log; không train hay forward test."""
from __future__ import annotations
import ast, hashlib, json, re, subprocess, sys
from pathlib import Path
from unittest.mock import patch
import importlib.metadata
import numpy as np
import pandas as pd
import nbformat
from openpyxl import load_workbook
from paths import CODE, SUB, DATA, REPO
from train import write_json
from freeze_final import assert_integrity, config_for
from test_once import validate_cached, export_all
from eval import load_group, read_pred, check_against_csv, compute_metrics
from build_workbook import source_tables, value, LABELS, SHEETS

def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def audit_submission():
    checks=[]
    def passed(name,evidence):checks.append({'criterion':name,'status':'Đạt','evidence':evidence})
    source_diff=subprocess.run(['git','diff','fe9791a','--','eval.py','starter','tests','README.md','GUIDE.md','RUBRIC.md'],cwd=REPO,text=True,capture_output=True,check=True).stdout
    if source_diff:raise AssertionError('Mã/tài liệu gốc bị sửa')
    passed('Giữ nguyên nguồn lớp học','git diff fe9791a: eval.py, starter/, tests/, README/ GUIDE/ RUBRIC không đổi')
    frozen=json.loads((SUB/'configs/frozen_final.json').read_text());assert_integrity(frozen)
    if SUB==CODE.parent:
        old=subprocess.run(['git','show','1629d9e:'+str((SUB/'configs/frozen_final.json').relative_to(REPO))],cwd=REPO,capture_output=True,check=True).stdout
        if hashlib.sha256(old).hexdigest()!=sha(SUB/'configs/frozen_final.json'):raise AssertionError('File chốt bị sửa sau commit')
    passed('Chốt trước test và giữ nguyên toán học','1629d9e; fingerprint train/model/dataset/loss/inference; ngoại lệ cờ runtime 6c07c9f trước test')
    gate5=json.loads((SUB/'evidence/stage5.json').read_text())
    assert gate5['split']['n']=={'train':10501,'val':3501,'test':3507}
    assert gate5['split']['union']==17509 and not any(gate5['split']['overlap'].values())
    passed('Fold0, EDA và sanity','stage5.json; lớp/ảnh mẫu/augmentation/overfit; giữ annotation conflict nguồn')
    b=pd.read_csv(SUB/'tables/Backbones.csv');t=pd.read_csv(SUB/'tables/Training.csv');inf=pd.read_csv(SUB/'tables/Inference.csv')
    assert len(b)==5 and len(t)==8 and len(inf)==9
    assert set(json.loads((SUB/'evidence/stage7.json').read_text())['axes'])=={'A','B','C'}
    for frame in [b,t]:
        for _,r in frame.iterrows():
            p=read_pred(SUB/'predictions'/f'{r.exp_id}_seed0_val.csv')
            check_against_csv(p,DATA/'labels/val_subset0.csv','val')
            m=compute_metrics(p.y_true,p.y_pred,p.probs)
            assert np.isclose(m['macro_f1'],r.macro_f1_val,atol=1e-12,rtol=0)
            assert np.isclose(m['top1'],r.top1_val,atol=1e-12,rtol=0)
    for _,r in inf.iterrows():
        p=read_pred(SUB/'predictions'/f'{r.exp_id}_seed0_val.csv');check_against_csv(p,DATA/'labels/val_subset0.csv','val')
        m=compute_metrics(p.y_true,p.y_pred,p.probs)
        for k in ['macro_f1','top1','ece','nll']:assert np.isclose(m[k],r[k+'_val'],atol=1e-12,rtol=0)
    passed('5 backbone, 3 trục, kết hợp và 9 cấu hình suy luận','Backbones/Training/Inference CSV khớp evaluator; các gate6/7/8; sàng lọc một seed')
    runs=[]
    for eid,seed in [(e,0) for e in b.exp_id]+[(e,0) for e in t.exp_id]+[(e,s) for e in ['F01','T00'] for s in frozen['seeds']]:
        folder=SUB/'logs'/eid/f'seed{seed}';cfg=json.loads((folder/'config.json').read_text())['config']
        s=json.loads((folder/'summary.json').read_text());h=pd.read_csv(folder/'history.csv')
        assert h.epoch.tolist()==list(range(1,11)) and s['epochs_completed']==10
        assert cfg['seed']==seed and cfg['epochs']==10 and cfg['batch_size']==32 and cfg['img_size']==128
        curve=SUB/'curves'/Path(s['curve']).name;assert curve.is_file()
        runs.append({'exp_id':eid,'seed':seed,'epochs':10,'curve':str(curve.relative_to(SUB))})
    assert len(runs)==19
    passed('19 lần train thật và 19 biểu đồ','Mỗi history đủ epoch1–10, mọi summary/curve có thật; không dùng pilot làm chung kết')
    freeze_sha=sha(SUB/'configs/frozen_final.json');fval=pd.read_csv(SUB/'tables/FinalValidation.csv')
    for eid in ['F01','T00']:
        for seed in frozen['seeds']:
            cfg=config_for(eid,seed);folder=SUB/'logs'/eid/f'seed{seed}'
            ledger=json.loads((folder/'test_ledger.json').read_text());s=json.loads((folder/'summary.json').read_text())
            validate_cached(ledger,cfg,freeze_sha,s['checkpoint_sha256'])
            assert ledger['n']==3507 and ledger['forward_dataset_passes']==1
            assert ledger['started_utc']>frozen['created_utc']
            v=json.loads((folder/'final_val.json').read_text())
            assert v['temperature']==ledger['temperature']
            if eid=='F01':assert ledger['temperature_fit_split']=='val'
            else:assert ledger['temperature_fit_split'] is None and ledger['temperature']==1.
            calibrated=read_pred(SUB/'predictions'/f'{eid}_seed{seed}_test.csv')
            raw=read_pred(SUB/'predictions'/f'{eid}_uncal_seed{seed}_test.csv')
            assert np.array_equal(calibrated.y_pred,raw.y_pred)
            vp=read_pred(SUB/'predictions'/f'{eid}_seed{seed}_val.csv');check_against_csv(vp,DATA/'labels/val_subset0.csv','val')
            row=fval[(fval.exp_id==eid)&(fval.seed==seed)].iloc[0]
            assert np.isclose(row.macro_f1_val,compute_metrics(vp.y_true,vp.y_pred,vp.probs)['macro_f1'],atol=1e-12,rtol=0)
    # Nếu cache vô tình đi tới model, kiểm tra này thất bại trước mọi forward.
    with patch('test_once.load_checkpoint',side_effect=AssertionError('Không được gọi model test lần hai')):
        cached=export_all()
    assert len(cached)==6
    passed('Test một lượt/seed và TS chỉ val','Sáu ledger độc quyền complete/pass_count1/n3507; argmax TS giữ nguyên; gọi lại chỉ xác minh cache, mock chặn model')
    gate=json.loads((SUB/'evidence/stage11.json').read_text());ft=pd.read_csv(SUB/'tables/Final.csv');pc=pd.read_csv(SUB/'tables/PerClass.csv')
    for eid in ['F01','T00']:
        g=load_group([str(SUB/'predictions'/f'{eid}_seed*_test.csv')],str(DATA/'labels/test_subset0.csv'))
        assert len(g.preds)==3
        for k in ['macro_f1','top1','ece','nll']:
            assert np.allclose(g.summary[k],[gate['groups'][eid][k]['mean'],gate['groups'][eid][k]['std']],atol=1e-12,rtol=0)
        for p,m in zip(g.preds,g.metrics):
            row=ft[(ft.exp_id==eid)&(ft.seed==p.seed)].iloc[0]
            for k in ['macro_f1','top1','ece','nll']:assert np.isclose(row[k+'_test'],m[k],atol=1e-12,rtol=0)
        rows=pc[pc.exp_id==eid]
        assert rows.test_count.sum()==3507
        for k in ['precision','recall','f1']:
            assert np.allclose(rows[k],g.summary[k][0],atol=1e-12,rtol=0)
            assert np.allclose(rows[k+'_std'],g.summary[k][1],atol=1e-12,rtol=0)
    passed('Metric mean/sample std và per-class','Tính lại từ đủ CSV test; đối chiếu Final/PerClass/stage11 và evaluator ddof=1')
    lat=pd.read_csv(SUB/'tables/Latency.csv');assert len(lat)==22 and (lat.n>=50).all() and (lat.warmups>=10).all()
    for path in list((SUB/'logs/inference').rglob('*latency*.json'))+list((SUB/'logs/F01/seed0').glob('*latency.json'))+list((SUB/'logs/T00/seed0').glob('*latency.json')):
        record=json.loads(path.read_text());assert record['n']>=50 and len(record['samples_ms'])==record['n']
        for k,q in [('p50',50),('p95',95),('p99',99)]:assert np.isclose(record[k],np.percentile(record['samples_ms'],q),atol=1e-12)
    passed('Latency thực và giới hạn phạm vi','22 hàng batch1/32; 10 warmup/50 mẫu sync; p50/p95/p99 từ mẫu thô; không tính CPU preprocessing')
    frames=source_tables();wb=load_workbook(SUB/'results.xlsx',data_only=True)
    assert wb.sheetnames==SHEETS
    for name,df in frames.items():
        ws=wb[name];assert ws.max_row==len(df)+5
        assert [ws.cell(5,j).value for j in range(1,len(df.columns)+1)]==[LABELS.get(c,c) for c in df.columns]
        for i,row in enumerate(df.itertuples(index=False,name=None),6):
            for j,v in enumerate(row,1):
                a=ws.cell(i,j).value;e=value(v)
                if isinstance(e,float):assert np.isclose(a,e,atol=1e-14,rtol=1e-14)
                else:assert a==e
        assert ws.freeze_panes and ws.auto_filter.ref
    passed('Workbook bảy sheet đúng từng ô','Đọc XLSX đã lưu, đối chiếu CSV và mean/std; tiêu đề Việt/đơn vị/filter/freeze; preview mọi sheet/cột đã xem')
    report=(SUB/'report.md').read_text();assert len(re.findall(r'^## [1-9]\. ',report,re.M))==9
    for link in re.findall(r'\]\(([^)]+)\)',report):
        if not link.startswith('https://'):assert (SUB/link).is_file(),link
    manifest=json.loads((SUB/'configs/report_sources.json').read_text())
    for path,digest in manifest.items():assert sha(SUB/path)==digest
    passed('Báo cáo tiếng Việt, giới hạn và lỗi thật','9 mục; số/ảnh từ nguồn có hash; nêu n3/fold0/budget/cache/head/MPS/runtime fix; liên kết nội bộ đủ')
    notebook=nbformat.read(CODE/'lab_day2.ipynb',as_version=4);nbformat.validate(notebook)
    for cell in notebook.cells:
        if cell.cell_type=='code':ast.parse(cell.source);assert not cell.outputs and cell.execution_count is None
    for path in CODE.rglob('*'):
        if path.suffix in ['.py','.ipynb']:
            unfinished=('TO'+'DO','Not'+'ImplementedError')
            assert not any(token in path.read_text() for token in unfinished),path
    passed('Notebook và mã hoàn thiện','nbformat/AST mọi code cell; không output cũ hoặc stub; cùng engine CLI; chưa chạy Colab toàn bộ')
    test_runs={}
    for label,folder in [('goc',REPO/'tests'),('bai_nop',CODE/'tests')]:
        result=subprocess.run([sys.executable,'-m','unittest','discover','-s',str(folder),'-v'],cwd=REPO,text=True,capture_output=True)
        output=result.stdout+result.stderr
        (SUB/'evidence'/f'final_tests_{label}.log').write_text('\n'.join(line.rstrip() for line in output.splitlines())+'\n')
        if result.returncode:raise AssertionError(output)
        count=int(re.search(r'Ran (\d+) tests',output).group(1));test_runs[label]=count
    assert test_runs=={'goc':38,'bai_nop':25}
    passed('Kiểm tra CPU thực thi','38 test gốc + 25 test bài nộp đạt; không thêm lượt test dataset')
    tracked=subprocess.run(['git','ls-files'],cwd=REPO,text=True,capture_output=True,check=True).stdout.splitlines()
    banned=[f for f in tracked if f.startswith(('data/','.venv/')) or Path(f).suffix in ['.pt','.pth','.safetensors','.npz','.zip']]
    assert not banned,banned
    passed('Không commit dữ liệu/trọng số','git ls-files không chứa data/.venv/checkpoint/safetensors/NPZ/ZIP; các sản phẩm nhỏ được commit')
    checklist=['# Kiểm toán cuối theo RUBRIC','',
        '| Mục | Kết quả | Bằng chứng |','|---|---|---|']
    checklist += [f"| {c['criterion']} | {c['status']} | {c['evidence']} |" for c in checks]
    checklist += ['','## Đối chiếu các mục chấm','',
        '- P1–P4: đủ sáu sản phẩm, nguồn thực, fold0, CSV chung kết và mốc đủ ba seed.',
        '- A/B/D/E/F/G/H: các bằng chứng kiểm tra trong bảng trên. Không tự gán điểm toàn phần A–H.',
        '- C: đủ ba trục, kiểm soát và T07 kết hợp LS+CutMix; CutMix đơn lẻ không tốt hơn mốc. Vì rubric nói kết hợp các yếu tố tốt, tiêu chí này cần giảng viên cân nhắc; không tuyên bố thử hai yếu tố đều tốt.',
        '- I: eval.py nguyên bản tự chấm 19/20. Hai recall tăng so T00 nhưng chưa vượt mốc bài báo. Giảng viên quyết định điểm.',
        '- Điểm thưởng: không yêu cầu hoặc tuyên bố có; chưa thực nghiệm thêm fold/DINO/distillation/ONNX/TTA thích ứng.',
        '- Notebook hợp lệ và gọi cùng module; chưa chạy toàn bộ trên Colab. Biểu đồ XLSX kiểm tra bằng preview đọc file đã lưu; không mở ứng dụng Excel native.',
        '- Test đã mở: cấu hình giữ nguyên, không huấn luyện lại hoặc chạy model test lần hai trong audit.']
    (SUB/'evidence/final_audit.md').write_text('\n'.join(checklist)+'\n')
    record={'passed':True,'checks':checks,'training_runs':runs,'tests':test_runs,'test_dataset_passes':6,
        'test_passes_per_seed':1,'official_I_points':19,'official_I_max':20,
        'rubric_combination_note':'T07 có LS tốt/CutMix không tốt: giảng viên cân nhắc điểm kết hợp yếu tố tốt',
        'notebook_full_colab_execution':False,'native_excel_check':False}
    write_json(SUB/'evidence/stage15.json',record)
    print(json.dumps({'passed':True,'checks':len(checks),'tests':test_runs,'test_forward_added':0},ensure_ascii=False))
    return record

if __name__=='__main__':audit_submission()
